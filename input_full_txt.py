from docx import Document

new_filename = input("ファイル名: ")
draft_content = ["１行目", "２行目"]

document = Document("TEST用起案文書様式.docx")

table = document.tables[0]

line_num = 6

for draft in draft_content:
    # 表のline_num行目、0列目にあるセルを取得する
    print(draft)
    cell = table.cell(line_num, 0)
    # セル内のテキストを上書きする
    cell.text = draft
    line_num += 1


document.save(f"{new_filename}.docx")
